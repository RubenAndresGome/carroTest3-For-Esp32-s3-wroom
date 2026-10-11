import docx, re, sys
sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r"D:\ITESS\TICS\VERANO 2026\TEST_RESIDENCIAS\AnteProyecto\documentos_corregidos\informe\informe_Ruben_Andres_Gomez_Molina_alfa_0_6_2.docx"

doc = docx.Document(DOC_PATH)
body = doc._body._body

print("="*60)
print("AUDITORÍA FORENSE DEL INFORME ALFA 0.6.2")
print("="*60)

# 1. Portada
xml_str = doc._element.xml
print("\n[1] VERSIÓN EN PORTADA:")
if ". Alfa 0.6.2" in xml_str:
    print("   ✓ EXITOSO: Portada actualizada a '. Alfa 0.6.2'")
else:
    print("   ✗ ERROR: No se encontró '. Alfa 0.6.2' en portada!")

# 2. Resumen
print("\n[2] RESUMEN EN ESPAÑOL:")
print("   P17:", doc.paragraphs[17].text[:120], "...")
print("   P18:", doc.paragraphs[18].text[:120], "...")
raw_dirs = any('desktop_app/' in doc.paragraphs[k].text for k in range(15, 25))
if not raw_dirs:
    print("   ✓ EXITOSO: Lista cruda de carpetas eliminada del Resumen.")
else:
    print("   ✗ ERROR: Todavía hay menciones de carpetas crudas en el Resumen!")

# 3. Abstract
print("\n[3] ABSTRACT EN INGLÉS:")
print("   P47:", doc.paragraphs[47].text[:140], "...")
if "step-command execution success rate of 79.7%" in doc.paragraphs[47].text:
    print("   ✓ EXITOSO: Abstract corregido con 'step-command execution success rate of 79.7%'")
else:
    print("   ✗ ERROR: Frase de 79.7% no coincide exactamente.")

if "overall command execution success rate of 79.7%" not in xml_str:
    print("   ✓ EXITOSO: Frase ambigua 'overall command execution success rate of 79.7%' erradicada.")

# 4. Recomendaciones
print("\n[4] RECOMENDACIONES (Alfa 0.01):")
if "Alfa 0.01" not in xml_str and "alfa 0.01" not in xml_str.lower():
    print("   ✓ EXITOSO: 'Alfa 0.01' completamente erradicada.")
    print("   P1050 actual:", doc.paragraphs[1050].text[:120], "...")
else:
    print("   ✗ ERROR: Aún existe mención de Alfa 0.01 en el documento!")

# 5. Anexo D y Anexo E
print("\n[5] ANEXOS:")
p1390_txt = doc.paragraphs[1390].text
print("   P1390 (Anexo D):", repr(p1390_txt))
if not p1390_txt.endswith('['):
    print("   ✓ EXITOSO: Carácter '[' eliminado de Anexo D.")
else:
    print("   ✗ ERROR: Anexo D todavía tiene '[' al final!")

p1397_txt = doc.paragraphs[1397].text
print("   P1397 (Anexo E):", repr(p1397_txt))

# 6. Tabla de Ilustraciones / Figuras
print("\n[6] TABLA DE FIGURAS EN EL FRONT MATTER:")
cat_idx = None
ecuaciones_idx = None
for i, child in enumerate(body[:120]):
    txt = ''.join(child.itertext()).strip().lower()
    if cat_idx is None and ('catálogo de ilustraciones' in txt or 'catalogo de ilustraciones' in txt):
        cat_idx = i
    elif cat_idx is not None and ('índice de ecuaciones' in txt or 'indice de ecuaciones' in txt):
        ecuaciones_idx = i
        break

fig_lines = []
if cat_idx and ecuaciones_idx:
    for k in range(cat_idx + 1, ecuaciones_idx):
        t = ''.join(body[k].itertext()).strip()
        fig_lines.append(t)

print(f"   Total figuras indexadas: {len(fig_lines)}")
pages = []
for idx, f in enumerate(fig_lines, 1):
    parts = f.split('\t')
    title = parts[0]
    p_num = parts[1] if len(parts) > 1 else 'N/A'
    # clean p_num
    clean_p = ''.join([c for c in p_num if c.isdigit()])
    print(f"   {idx:02d}. {title[:55]} -> Pág {clean_p or p_num}")
    if clean_p.isdigit():
        pages.append(int(clean_p))

# Verificar monotonicidad de páginas
if pages:
    is_monotonic = all(pages[i] <= pages[i+1] for i in range(len(pages)-1))
    if is_monotonic:
        print("   ✓ EXITOSO: Todas las páginas de figuras aumentan monótonamente (sin saltos hacia atrás).")
    else:
        print("   ! ATENCIÓN: Revisar monotonicidad de páginas:")
        for i in range(len(pages)-1):
            if pages[i] > pages[i+1]:
                print(f"      Salto: Fig {i+1} (Pág {pages[i]}) > Fig {i+2} (Pág {pages[i+1]})")

print("\n" + "="*60)
print("AUDITORÍA FINALIZADA")
print("="*60)
