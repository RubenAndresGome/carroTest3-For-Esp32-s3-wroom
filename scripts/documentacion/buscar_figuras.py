import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'D:\ITESS\TICS\VERANO 2026\TEST_RESIDENCIAS\AnteProyecto\documentos_corregidos\informe\informe_Ruben_Andres_Gomez_Molina_alfa_0_6_1.docx')

queries = [
    'Figura 1 Instituto',
    'Simplificación de mallas',
    'Figura 4 Diagrama',
    'Figura 5 Extracto',
    'Figura 6 Ensamble',
    'Figura 8 https',
    'Figura 12 Mapa',
    'Figura 13 Mapa',
    'Figura 14 Mapa',
    'Figura 15 Zona',
    'Figura 16 Zona',
    'Figura 17 Diagrama',
    'Skill para actualizar',
    'Figura 20 Robot S3',
    'Figura 21 Ensamble',
    'Figura 22 Primera',
    'Figura 23 Montaje',
    'Figura 24 Primeras',
    'Figura 25 Sistema',
    'Figura 26 Cableado',
    'Figura 27 Evidencia',
    'Figura 28 herramientas',
    'Figura 29 Pista',
    'Figura 30 Videos',
    'Figura 31 Aplicación',
    'Figura 32 Basado',
    'Figura 33 Fuente',
    'Figura 35 Carrito',
    'Carta de Aceptación',
    'Carta de liberación',
    'Figura 38 Esquema',
    'Collage de Evidencias',
    'Figura 40 ANEXO',
    'Figura 41 Complicación',
    'Figura 42 Uso',
    'Figura 43 Visual',
    'Figura 44 Diversas',
]

for idx, q in enumerate(queries, 1):
    found = False
    for pi, p in enumerate(doc.paragraphs[98:], start=98):
        if q.lower() in p.text.lower():
            print(f'{idx:02d}. Query "{q[:22]}" -> P{pi}: {p.text.strip()[:65]}')
            found = True
            break
    if not found:
        for ti, t in enumerate(doc.tables):
            for ri, r in enumerate(t.rows):
                for ci, c in enumerate(r.cells):
                    if q.lower() in c.text.lower():
                        clean_txt = c.text.strip().replace('\n', ' ')[:65]
                        print(f'{idx:02d}. Query "{q[:22]}" -> T{ti}[{ri},{ci}]: {clean_txt}')
                        found = True
                        break
                if found: break
            if found: break
    if not found:
        print(f'{idx:02d}. Query "{q[:22]}" -> NOT FOUND!')
