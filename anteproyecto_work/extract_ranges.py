from pathlib import Path
import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(r"D:\ITESS\TICS\VERANO 2026\TESTs\AnteProyecto")
targets = [
    ("Reporte_Final_Robot_Movil_RubenAndresGomezMolina.docx", 0, 103),
    ("Anteproyecto-copy rescue.docx", 195, 252),
    ("Anteproyecto-copy rescue.docx", 630, 669),
]

for filename, start, end in targets:
    doc = Document(ROOT / filename)
    print(f"\n{'=' * 88}\n{filename} | P{start + 1}–P{end + 1}")
    for n in range(start, min(end + 1, len(doc.paragraphs))):
        p = doc.paragraphs[n]
        text = " ".join(p.text.split())
        if text:
            print(f"P{n + 1:04d} [{p.style.name}] {text}")
