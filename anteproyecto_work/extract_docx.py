from pathlib import Path
import sys
from docx import Document

SOURCE = Path(r"D:\ITESS\TICS\VERANO 2026\TESTs\AnteProyecto")
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

for path in sorted(SOURCE.glob("*.docx")):
    print(f"\n{'=' * 88}\nARCHIVO: {path.name}")
    doc = Document(path)
    print(f"PARRAFOS: {len(doc.paragraphs)} | TABLAS: {len(doc.tables)}")
    print("\n[ENCABEZADOS Y PÁRRAFOS RELEVANTES]")
    for i, p in enumerate(doc.paragraphs):
        text = " ".join(p.text.split())
        if not text:
            continue
        style = p.style.name if p.style else ""
        if "Heading" in style or "Título" in style or "Title" in style or i < 35:
            print(f"P{i + 1:04d} [{style}] {text[:500]}")
    print("\n[TABLAS]")
    for ti, table in enumerate(doc.tables, 1):
        print(f"TABLA {ti} ({len(table.rows)} x {len(table.columns)})")
        for row in table.rows[:12]:
            cells = [" ".join(c.text.split())[:240] for c in row.cells]
            print(" | ".join(cells))
        if len(table.rows) > 12:
            print("...")
